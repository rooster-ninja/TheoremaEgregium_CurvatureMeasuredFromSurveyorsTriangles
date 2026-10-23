import json,math
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
RAD=math.pi/180/3600
rows=json.load(open('rows_t.json')); sig=json.load(open('rule_results.json'))['head']['sig']; ch=json.load(open('chain_results.json'))
inchain={n for sv in ('US 39th parallel','US Eastern Oblique','Maclear, Cape 1837–48') for n in ch[sv]['names']}
F='Arial'; BLUE=Font(name=F,color='0000FF'); BLACK=Font(name=F); BOLD=Font(name=F,bold=True); HDR=Font(name=F,bold=True,color='FFFFFF')
HF=PatternFill('solid',fgColor='2A5FA8'); YEL=PatternFill('solid',fgColor='FFFF00'); GRN=Font(name=F,color='008000')
wb=Workbook()
# ---------------- Inputs
ws=wb.active; ws.title='Inputs'
ws['A1']='Curvature measured from surveyors\' triangles · workbook'; ws['A1'].font=Font(name=F,bold=True,size=13)
ws['A2']='Blue = input you can change · black = formula · green = link to another sheet. Change the SNR cut or tier weights and the Result sheet recalculates.'; ws['A2'].font=Font(name=F,italic=True)
inp=[('R0, fixed scale radius for SNR (m)',6371000,'Fixed before the run; not a fit (RULE.md)'),
     ('Arc-seconds per radian',206264.806247,'360·3600/(2π)'),
     ('SNR cut (keep triangles with SNR ≥ cut)',2,'Declared rule: 2. Try 0, 1 or 3 (section 4 sensitivity)'),
     ('Tier A weight',1.0,'Observed angles + 3 printed sides'),('Tier B weight',0.5,'Observed angles + 1–2 printed sides'),('Tier C weight',0.25,'Observed angles, coordinate area only')]
for i,(k,v,n) in enumerate(inp):
    r=4+i; ws.cell(r,1,k).font=BLACK; c=ws.cell(r,2,v); c.font=BLUE; c.fill=YEL if 'cut' in k or 'weight' in k else PatternFill(); ws.cell(r,3,n).font=Font(name=F,italic=True,color='5A5650')
ws.column_dimensions['A'].width=46; ws.column_dimensions['B'].width=16; ws.column_dimensions['C'].width=60
# named refs
R0='Inputs!$B$4'; ASR='Inputs!$B$5'; CUT='Inputs!$B$6'; WA='Inputs!$B$7'; WB='Inputs!$B$8'; WC='Inputs!$B$9'
# ---------------- Triangles
t=wb.create_sheet('Triangles')
hd=['Survey','Triangle','Tier','Area used (km²)','Observed excess Σangles−180° (″)','K ellipsoid 1/(MN) (″ per km²)','Survey closure σ (″)','Expected excess A·K (″)','SNR','Tier weight','In result','In chain sum','Area from coordinates (km²)','Printed excess (″)','QC note','Fit weight w = tier/σ² (if in result)','Residual² (if in result)']
for j,h in enumerate(hd,1):
    c=t.cell(1,j,h); c.font=HDR; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='top')
r=2
for x in rows:
    A=(x['A_side'] if (x.get('A_side') and x['tier']!='C') else x['area'])/1e6
    t.cell(r,1,x['survey']); t.cell(r,2,x['name']); t.cell(r,3,x['tier'] or '')
    t.cell(r,4,round(A,4)).font=BLUE
    e=t.cell(r,5,None if x['E'] is None else round(x['E'],4)); e.font=BLUE
    t.cell(r,6,x['Kref']*1e6/RAD).font=BLUE
    s=sig.get(x['survey']); t.cell(r,7,None if s is None else round(s,4)).font=BLUE
    t.cell(r,8,f'=D{r}*F{r}')
    t.cell(r,9,f'=IF(G{r}="","",D{r}*1E6/{R0}^2*{ASR}/G{r})')
    t.cell(r,10,f'=IF(C{r}="A",{WA},IF(C{r}="B",{WB},IF(C{r}="C",{WC},0)))')
    t.cell(r,11,f'=IF(AND(E{r}<>"",G{r}<>"",J{r}>0),IF(I{r}>={CUT},1,0),0)')
    t.cell(r,12,1 if x['name'] in inchain and x['survey'] in ch else 0).font=BLUE
    t.cell(r,13,round(x['area']/1e6,4)).font=BLUE
    t.cell(r,14,x['Ep']).font=BLUE; t.cell(r,15,x['qc'])
    t.cell(r,16,f'=IF(K{r}=1,J{r}/G{r}^2,0)')
    t.cell(r,17,f'=IF(K{r}=1,((E{r}-Result!$B$7*D{r})/G{r})^2,0)')
    for j in (1,2,3,15): t.cell(r,j).font=BLACK
    r+=1
N=r-1
for j,w in enumerate([24,52,6,14,16,16,12,14,8,9,9,9,14,12,60,16,14],1): t.column_dimensions[L(j)].width=w
t.freeze_panes='C2'
t['D1'].comment=Comment('From printed side lengths (Heron, or one side + observed angles) for tiers A/B; from WGS84 coordinates for tier C. Source: arcs_parallel_master.kml; computed in scripts/fit2.py.','analysis')
t['E1'].comment=Comment('Sum of the three observed angles minus 180°, from the printed tables as transcribed in the KML. Three entries corrected from Schott misprints (see Source_check).','analysis')
t['G1'].comment=Comment('1.4826 × median absolute deviation of each survey\'s residuals from its own one-parameter fit (RULE.md).','analysis')
# ---------------- Result
s=wb.create_sheet('Result',1)
rng=lambda c:f'Triangles!${c}$2:${c}${N}'
A,E,K,S,W,IN=rng('D'),rng('E'),rng('F'),rng('G'),rng('J'),rng('K'); P,Q=rng('P'),rng('Q')
items=[('Triangles in the result',f'=SUM({IN})',None),
 ('Σ w·A·ε',f'=SUMPRODUCT({P},{A},{E})',None),
 ('Σ w·A²',f'=SUMPRODUCT({P},{A},{A})',None),
 ('K measured (″ per km²)','=B5/B6',None),
 ('K ellipsoid, same weights (″ per km²)',f'=SUMPRODUCT({P},{A},{K},{A})/B6',None),
 ('Ratio q = K / K ellipsoid','=B7/B8','Main result'),
 ('Variance of K (sandwich)',f'=SUMPRODUCT({P},{P},{A},{A},{S},{S})/B6^2',None),
 ('χ²/dof',f'=SUM({Q})/(B4-1)',None),
 ('σ of K (″ per km²), inflated if χ²/dof > 1','=SQRT(B10)*MAX(1,SQRT(B11))',None),
 ('σ of q','=B12/B8','Main result'),
 ('z (q vs 1)','=(B9-1)/B13',None),
 ('K measured (m⁻², ×10⁻¹⁴)',f'=B7/{ASR}/1E6*1E14','Main result'),
 ('σ of K (m⁻², ×10⁻¹⁴)',f'=B12/{ASR}/1E6*1E14',None),
 ('Radius 1/√K (km)',f'=1/SQRT(B15*1E-14)/1000','Main result')]
s['A1']='Main result · recalculates from the Triangles sheet and the Inputs'; s['A1'].font=Font(name=F,bold=True,size=13)
s['A2']='Expected with the declared inputs: 391 triangles, q = 0.9953 ± 0.0048, K = 2.4501 × 10⁻¹⁴ m⁻², R = 6,389 km (scripts/rule_fit.py).'; s['A2'].font=Font(name=F,italic=True)
for i,(k,f,tag) in enumerate(items):
    r=4+i; s.cell(r,1,k).font=BOLD if tag else BLACK; c=s.cell(r,2,f); c.font=Font(name=F,bold=bool(tag)); s.cell(r,3,tag or '').font=Font(name=F,italic=True,color='5A5650')
for r in (7,8): s.cell(r,2).number_format='0.000000'
s['B9'].number_format='0.0000'; s['B13'].number_format='0.0000'; s['B14'].number_format='0.00'; s['B11'].number_format='0.00'
s['B15'].number_format='0.0000'; s['B16'].number_format='0.0000'; s['B17'].number_format='#,##0'
s.column_dimensions['A'].width=46; s.column_dimensions['B'].width=18; s.column_dimensions['C'].width=16
# per survey
s['A20']='By survey (unweighted within a survey, as in section 4)'; s['A20'].font=BOLD
for j,h in enumerate(['Survey','Triangles','K (″/km²)','K ellipsoid (″/km²)','q','σ of q'],1):
    c=s.cell(21,j,h); c.font=HDR; c.fill=HF
SV=rng('A'); SG=rng('G')
for i,sv in enumerate(['US 39th parallel','US Eastern Oblique','Maclear, Cape 1837–48','Europe 52°N (Helmert)','Transvaal–Rhodesia','India GTS meridional','England OS 1787–96']):
    r=22+i; s.cell(r,1,sv).font=BLUE
    s.cell(r,2,f'=SUMPRODUCT(({SV}=A{r})*{IN})')
    s.cell(r,3,f'=SUMPRODUCT(({SV}=A{r})*{IN},{A},{E})/SUMPRODUCT(({SV}=A{r})*{IN},{A},{A})')
    s.cell(r,4,f'=SUMPRODUCT(({SV}=A{r})*{IN},{A},{K},{A})/SUMPRODUCT(({SV}=A{r})*{IN},{A},{A})')
    s.cell(r,5,f'=C{r}/D{r}').number_format='0.000'
    s.cell(r,6,f'=INDEX({SG},MATCH(A{r},{SV},0))/SQRT(SUMPRODUCT(({SV}=A{r})*{IN},{A},{A}))/D{r}').number_format='0.000'
    for j in (3,4): s.cell(r,j).number_format='0.000000'
s['A30']='σ of q per survey here is before χ² inflation; section 4 of the report applies it.'; s['A30'].font=Font(name=F,italic=True,color='5A5650')
for j,w in enumerate([46,18,16,18,10,10],1): s.column_dimensions[L(j)].width=w
# ---------------- Chain sums
c=wb.create_sheet('Chain_sums')
c['A1']='Whole-chain sums over non-overlapping triangles (Appendix B)'; c['A1'].font=Font(name=F,bold=True,size=13)
for j,h in enumerate(['Chain','Triangles','Area (km²)','Σ excess observed (″)','Σ expected (″)','Noise σ√n (″)','K/K ellipsoid','Shortfall per triangle (″)'],1):
    x=c.cell(3,j,h); x.font=HDR; x.fill=HF
CI=rng('L'); EX=rng('H')
for i,sv in enumerate(['US 39th parallel','US Eastern Oblique','Maclear, Cape 1837–48']):
    r=4+i; c.cell(r,1,sv).font=BLUE
    c.cell(r,2,f'=SUMPRODUCT(({SV}=A{r})*{CI})')
    c.cell(r,3,f'=SUMPRODUCT(({SV}=A{r})*{CI},{A})').number_format='#,##0'
    c.cell(r,4,f'=SUMPRODUCT(({SV}=A{r})*{CI},{E})').number_format='#,##0.0'
    c.cell(r,5,f'=SUMPRODUCT(({SV}=A{r})*{CI},{EX})').number_format='#,##0.0'
    c.cell(r,6,f'=INDEX({SG},MATCH(A{r},{SV},0))*SQRT(B{r})').number_format='0.0'
    c.cell(r,7,f'=D{r}/E{r}').number_format='0.000'
    c.cell(r,8,f'=(D{r}-E{r})/B{r}').number_format='+0.00;-0.00'
r=7; c.cell(r,1,'Core three combined').font=BOLD
for j,col in ((2,'B'),(3,'C'),(4,'D'),(5,'E')): c.cell(r,j,f'=SUM({col}4:{col}6)')
c.cell(r,6,'=SQRT(SUMPRODUCT(F4:F6,F4:F6))').number_format='0.0'; c.cell(r,7,'=D7/E7').number_format='0.000'; c.cell(r,8,'=(D7-E7)/B7').number_format='+0.00;-0.00'
c['C7'].number_format='#,##0'; c['D7'].number_format='#,##0.0'; c['E7'].number_format='#,##0.0'
for j,w in enumerate([26,10,12,18,14,13,13,20],1): c.column_dimensions[L(j)].width=w
# ---------------- Source check
v=wb.create_sheet('Source_check')
v['A1']='KML entries checked against the printed tables (Schott 1900, 1902)'; v['A1'].font=Font(name=F,bold=True,size=13)
v['A3']='Pages read entry by entry'; v['A3'].font=BOLD
for j,h in enumerate(['Report','Printed page','PDF page','Triangles on page','Angles','All match KML'],1):
    x=v.cell(4,j,h); x.font=HDR; x.fill=HF
pages=[('Schott 1900',414,424,9),('Schott 1900',474,485,9),('Schott 1900',588,616,9),('Schott 1900',591,619,9),('Schott 1902',131,153,12)]
for i,(rep,pp,pdf,n) in enumerate(pages):
    r=5+i
    for j,val in enumerate([rep,pp,pdf,n],1): v.cell(r,j,val).font=BLUE
    v.cell(r,5,f'=D{r}*3'); v.cell(r,6,'Yes').font=BLUE
v.cell(10,1,'Total').font=BOLD; v.cell(10,4,'=SUM(D5:D9)'); v.cell(10,5,'=SUM(E5:E9)')
v['A12']='Three printed misprints, corrected from the printed sides by the law of cosines'; v['A12'].font=BOLD
hd=['Triangle','Source','Station with misprint','Printed angle (°)','Side opposite it (m)','Other side b (m)','Other side c (m)','Angle from sides (°)','Corrected angle (°)']
for j,h in enumerate(hd,1):
    x=v.cell(13,j,h); x.font=HDR; x.fill=HF; x.alignment=Alignment(wrap_text=True)
mis=[('Mount Nebo – Wheeler Peak – Ibepah','Schott 1900 p. 591 no. 33','Ibepah','10 09 30.489',237769.08,99493.31,184375.27,'110 09 30.489'),
     ('Cherry Ridge – Twin Creek – Round Top','Schott 1900 p. 414 no. 25','Cherry Ridge','20 10 13.41',11997.12,24389.25,18130.38,'28 10 13.41'),
     ('Gunstock – Mount Pleasant – Mount Independence','Schott 1902 p. 131 no. 16','Mount Pleasant','19 54 04.38',88811.79,50110.16,71681.66,'91 54 04.38')]
for i,m in enumerate(mis):
    r=14+i
    for j,val in enumerate(m[:7],1): v.cell(r,j,val).font=BLUE
    v.cell(r,8,f'=DEGREES(ACOS((F{r}^2+G{r}^2-E{r}^2)/(2*F{r}*G{r})))').number_format='0.000'
    v.cell(r,9,m[7]).font=BLUE
v['A18']='Sides are the printed "Distances in metres" for each triangle. The angle from sides matches the corrected degrees to within the rounding of the printed sides.'; v['A18'].font=Font(name=F,italic=True,color='5A5650')
for j,w in enumerate([44,26,18,16,18,16,16,18,18],1): v.column_dimensions[L(j)].width=w
for wsx in wb.worksheets:
    for row in wsx.iter_rows():
        for cell in row:
            if cell.font is None or cell.font.name!=F: cell.font=Font(name=F,bold=cell.font.bold if cell.font else False,color=cell.font.color if cell.font else None,italic=cell.font.italic if cell.font else False,size=cell.font.size if cell.font else 11)
wb.save('/tmp/claude-0/-home-claude/029738cd-c575-57a3-afdb-e16d394ffc0b/scratchpad/pkg/Curvature_from_triangles.xlsx'); print('rows',N)
